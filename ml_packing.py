"""Experimental ML helpers for simple window and door packing.

The models are trained on deterministic synthetic examples generated from the
same documented constraints used by the demo.  Their predictions are shown for
comparison only.  Safety limits and the final packing result always come from
``calculate_rule_result``.
"""

from __future__ import annotations

import math
from dataclasses import dataclass
from io import BytesIO
from typing import Dict, Iterable, List, Tuple

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sklearn.compose import ColumnTransformer
from sklearn.ensemble import RandomForestClassifier, RandomForestRegressor
from sklearn.metrics import accuracy_score, mean_absolute_error, recall_score
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder


MAX_PALLET_WEIGHT_KG = 1000.0
PALLET_DEPTH_MM = 1200
PALLET_SIDE_DEPTH_MM = 550
MAX_PALLET_LENGTH_MM = 6800
MAX_PRODUCT_HEIGHT_MM = 2700
PALLET_HEIGHT_ALLOWANCE_MM = 200

INPUT_COLUMNS = [
    "Item",
    "Type",
    "Configuration",
    "Width (mm)",
    "Height (mm)",
    "Qty",
    "Frame weight (kg)",
    "Glass weight (kg)",
    "Packing thickness (mm)",
    "Handles installed",
    "Glass mode",
]

MODEL_FEATURES = [
    "Type",
    "Configuration",
    "Width (mm)",
    "Height (mm)",
    "Qty",
    "Frame weight (kg)",
    "Glass weight (kg)",
    "Packing thickness (mm)",
    "Handles installed",
    "Glass mode",
]

CATEGORICAL_FEATURES = [
    "Type",
    "Configuration",
    "Handles installed",
    "Glass mode",
]
NUMERIC_FEATURES = [c for c in MODEL_FEATURES if c not in CATEGORICAL_FEATURES]


@dataclass
class ModelBundle:
    units_model: Pipeline
    pallets_model: Pipeline
    ldm_model: Pipeline
    review_model: Pipeline
    metrics: Dict[str, float]
    training_rows: int
    test_rows: int


def _number(value, default: float = 0.0) -> float:
    try:
        if pd.isna(value):
            return default
        return float(str(value).replace("\xa0", "").replace(" ", "").replace(",", "."))
    except (TypeError, ValueError):
        return default


def normalize_type(value: object) -> str:
    text = str(value or "").strip().lower()
    if "window" in text or "lang" in text:
        return "Window"
    if "door" in text or "dur" in text:
        return "Door"
    return ""


def normalize_configuration(item_type: str, value: object) -> str:
    text = str(value or "").strip().lower()
    if item_type == "Window":
        if "fixed" in text or "fiks" in text:
            return "Fixed"
        if "mix" in text:
            return "Mixed"
        return "Openable"
    if "side" in text:
        return "Door + sidelight"
    if "double" in text or "2" == text:
        return "Double"
    return "Single"


def normalize_input(df: pd.DataFrame) -> pd.DataFrame:
    """Return the stable, system-neutral schema used by the ML demo."""
    source = df.copy()
    aliases = {
        "Item name": "Item",
        "Width mm": "Width (mm)",
        "Height mm": "Height (mm)",
        "Quantity": "Qty",
        "Frame kg/item": "Frame weight (kg)",
        "Frame kg": "Frame weight (kg)",
        "Unit weight (kg)": "Frame weight (kg)",
        "Glass kg/item": "Glass weight (kg)",
        "Glass kg": "Glass weight (kg)",
        "Packing thickness mm": "Packing thickness (mm)",
        "Handles": "Handles installed",
    }
    source = source.rename(columns={k: v for k, v in aliases.items() if k in source.columns})

    for column in INPUT_COLUMNS:
        if column not in source.columns:
            source[column] = None

    rows = []
    for index, row in source.iterrows():
        item = str(row.get("Item") or "").strip()
        if not item or item.lower() == "nan":
            item = f"Item {index + 1}"
        item_type = normalize_type(row.get("Type"))
        config = normalize_configuration(item_type, row.get("Configuration") or row.get("Type"))
        glass_mode = str(row.get("Glass mode") or "Glazed").strip().title()
        if glass_mode not in {"Glazed", "Unglazed", "Without Glass"}:
            glass_mode = "Glazed"
        if glass_mode == "Without Glass":
            glass_mode = "Without glass"
        handles = str(row.get("Handles installed") or "No").strip().title()
        handles = "Yes" if handles in {"Yes", "Y", "True", "1"} else "No"
        rows.append({
            "Item": item,
            "Type": item_type,
            "Configuration": config,
            "Width (mm)": _number(row.get("Width (mm)")),
            "Height (mm)": _number(row.get("Height (mm)")),
            "Qty": max(1, int(_number(row.get("Qty"), 1))),
            "Frame weight (kg)": _number(row.get("Frame weight (kg)")),
            "Glass weight (kg)": _number(row.get("Glass weight (kg)")),
            "Packing thickness (mm)": _number(row.get("Packing thickness (mm)"), 75),
            "Handles installed": handles,
            "Glass mode": glass_mode,
        })
    return pd.DataFrame(rows, columns=INPUT_COLUMNS)


def validate_input(df: pd.DataFrame) -> List[str]:
    errors: List[str] = []
    for index, row in df.iterrows():
        label = str(row.get("Item") or f"Row {index + 1}")
        if row.get("Type") not in {"Window", "Door"}:
            errors.append(f"{label}: Type must be Window or Door")
        for field in ("Width (mm)", "Height (mm)", "Qty", "Frame weight (kg)", "Packing thickness (mm)"):
            if _number(row.get(field)) <= 0:
                errors.append(f"{label}: {field} must be greater than 0")
        if row.get("Glass mode") in {"Glazed", "Unglazed"} and _number(row.get("Glass weight (kg)")) <= 0:
            errors.append(f"{label}: Glass weight is required for {row.get('Glass mode')}")
    return errors


def _count_limit(item_type: str, configuration: str) -> int:
    if item_type == "Door":
        return 6
    return {"Fixed": 10, "Openable": 6, "Mixed": 8}.get(configuration, 6)


def calculate_rule_result(row: pd.Series | Dict[str, object]) -> Dict[str, object]:
    r = dict(row)
    item_type = str(r["Type"])
    configuration = str(r["Configuration"])
    width = _number(r["Width (mm)"])
    height = _number(r["Height (mm)"])
    qty = max(1, int(_number(r["Qty"], 1)))
    frame = _number(r["Frame weight (kg)"])
    glass = _number(r["Glass weight (kg)"])
    thickness = max(1.0, _number(r["Packing thickness (mm)"], 75))
    glass_mode = str(r.get("Glass mode", "Glazed"))
    pallet_unit_weight = frame + glass if glass_mode == "Glazed" else frame

    count_limit = _count_limit(item_type, configuration)
    physical_capacity = max(1, math.floor(PALLET_SIDE_DEPTH_MM / thickness) * 2)
    weight_capacity = max(1, math.floor(MAX_PALLET_WEIGHT_KG / pallet_unit_weight)) if pallet_unit_weight > 0 else 1
    units_per_pallet = max(1, min(count_limit, physical_capacity, weight_capacity))
    pallets = math.ceil(qty / units_per_pallet)

    base_length = width + (100 if width <= 3000 else 200)
    handles_extra = (
        item_type == "Door"
        and str(r.get("Handles installed", "No")) == "Yes"
        and units_per_pallet > 4
    )
    pallet_length = width + 500 if handles_extra else base_length
    pallet_height = height + PALLET_HEIGHT_ALLOWANCE_MM
    total_ldm = round(pallets * pallet_length / 2000.0, 3)

    reasons = []
    if height > MAX_PRODUCT_HEIGHT_MM:
        reasons.append("Height over 2700 mm")
    if pallet_length > MAX_PALLET_LENGTH_MM:
        reasons.append("Pallet length over 6800 mm")
    if pallet_unit_weight > MAX_PALLET_WEIGHT_KG:
        reasons.append("One item over 1000 kg")

    return {
        "Rule units/pallet": units_per_pallet,
        "Rule pallets": pallets,
        "Rule pallet length (mm)": round(pallet_length),
        "Rule pallet height (mm)": round(pallet_height),
        "Rule total LDM": total_ldm,
        "Rule manual review": "Yes" if reasons else "No",
        "Rule reason": "; ".join(reasons) if reasons else "Within current limits",
        "Unit pallet weight (kg)": round(pallet_unit_weight, 2),
        "Physical capacity": physical_capacity,
        "Weight capacity": weight_capacity,
        "Count limit": count_limit,
    }


def build_pallet_rows(input_row: Dict[str, object], result: Dict[str, object]) -> List[Dict[str, object]]:
    qty = int(input_row["Qty"])
    capacity = int(result["Rule units/pallet"])
    unit_weight = float(result["Unit pallet weight (kg)"])
    rows = []
    remaining = qty
    for pallet_no in range(1, int(result["Rule pallets"]) + 1):
        units = min(capacity, remaining)
        side_a = math.ceil(units / 2)
        side_b = units - side_a
        rows.append({
            "Item": input_row["Item"],
            "Pallet": pallet_no,
            "Units": units,
            "Side A": side_a,
            "Side B": side_b,
            "Pallet weight (kg)": round(units * unit_weight, 2),
            "Pallet length (mm)": result["Rule pallet length (mm)"],
            "Pallet depth (mm)": PALLET_DEPTH_MM,
            "Pallet LDM": round(float(result["Rule pallet length (mm)"]) / 2000.0, 3),
            "Manual review": result["Rule manual review"],
        })
        remaining -= units
    return rows


def generate_training_data(size: int = 400, seed: int = 20260917) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    generated = []
    for index in range(size):
        item_type = "Window" if rng.random() < 0.62 else "Door"
        oversized = rng.random() < 0.08
        if item_type == "Window":
            configuration = rng.choice(["Fixed", "Fixed", "Openable", "Openable", "Mixed"])
            width = int(round(rng.uniform(600, 3000) / 10) * 10)
            height = int(round(rng.uniform(600, 2600) / 10) * 10)
            perimeter = 2 * (width + height) / 1000
            frame = perimeter * (4.5 if configuration == "Fixed" else 6.2) + rng.uniform(3, 30)
            handles = "No" if configuration == "Fixed" else "Yes"
        else:
            configuration = str(rng.choice(["Single", "Single", "Double", "Door + sidelight"]))
            width = int(round(rng.uniform(850 if configuration == "Single" else 1500, 3500) / 10) * 10)
            height = int(round(rng.uniform(1950, 2700) / 10) * 10)
            perimeter = 2 * (width + height) / 1000
            frame = perimeter * 8 + rng.uniform(10, 45)
            handles = "Yes" if rng.random() < 0.88 else "No"
        if oversized:
            if rng.random() < 0.7:
                height = int(rng.integers(2720, 2901))
            else:
                width = int(rng.integers(6650, 6901))
        glass_rate = 37 if rng.random() < 0.42 else 27
        glass = max(10.0, width * height / 1_000_000 * glass_rate * (0.82 if item_type == "Window" else 0.72))
        row = {
            "Item": f"Synthetic {index + 1}",
            "Type": item_type,
            "Configuration": configuration,
            "Width (mm)": width,
            "Height (mm)": height,
            "Qty": int(rng.integers(1, 21)),
            "Frame weight (kg)": round(max(20.0, frame), 1),
            "Glass weight (kg)": round(glass, 1),
            "Packing thickness (mm)": 75,
            "Handles installed": handles,
            "Glass mode": "Glazed",
        }
        target = calculate_rule_result(row)
        generated.append({
            **row,
            "Target units/pallet": target["Rule units/pallet"],
            "Target pallets": target["Rule pallets"],
            "Target total LDM": target["Rule total LDM"],
            "Target manual review": target["Rule manual review"],
            "Split": "Test" if (index + 1) % 5 == 0 else "Train",
            "Data origin": "Synthetic",
            "Production confirmed": "No",
        })
    return pd.DataFrame(generated)


def _pipeline(estimator) -> Pipeline:
    preprocessor = ColumnTransformer([
        ("categorical", OneHotEncoder(handle_unknown="ignore"), CATEGORICAL_FEATURES),
        ("numeric", "passthrough", NUMERIC_FEATURES),
    ])
    return Pipeline([("preprocess", preprocessor), ("model", estimator)])


def train_models() -> ModelBundle:
    data = generate_training_data()
    train = data[data["Split"] == "Train"]
    test = data[data["Split"] == "Test"]
    x_train, x_test = train[MODEL_FEATURES], test[MODEL_FEATURES]

    units = _pipeline(RandomForestRegressor(n_estimators=180, max_depth=12, random_state=11, n_jobs=-1))
    pallets = _pipeline(RandomForestRegressor(n_estimators=180, max_depth=14, random_state=12, n_jobs=-1))
    ldm = _pipeline(RandomForestRegressor(n_estimators=180, max_depth=14, random_state=13, n_jobs=-1))
    review = _pipeline(RandomForestClassifier(n_estimators=220, max_depth=12, class_weight="balanced", random_state=14, n_jobs=-1))

    units.fit(x_train, train["Target units/pallet"])
    pallets.fit(x_train, train["Target pallets"])
    ldm.fit(x_train, train["Target total LDM"])
    review.fit(x_train, train["Target manual review"])

    units_pred = np.maximum(1, np.rint(units.predict(x_test))).astype(int)
    pallets_pred = np.maximum(1, np.rint(pallets.predict(x_test))).astype(int)
    review_pred = review.predict(x_test)
    metrics = {
        "Units exact accuracy": float(accuracy_score(test["Target units/pallet"], units_pred)),
        "Pallets exact accuracy": float(accuracy_score(test["Target pallets"], pallets_pred)),
        "Pallets MAE": float(mean_absolute_error(test["Target pallets"], pallets_pred)),
        "LDM MAE": float(mean_absolute_error(test["Target total LDM"], ldm.predict(x_test))),
        "Manual review recall": float(recall_score(test["Target manual review"], review_pred, pos_label="Yes", zero_division=0)),
    }
    return ModelBundle(units, pallets, ldm, review, metrics, len(train), len(test))


def predict_batch(df: pd.DataFrame, bundle: ModelBundle) -> Tuple[pd.DataFrame, pd.DataFrame]:
    normalized = normalize_input(df)
    features = normalized[MODEL_FEATURES]
    ml_units = np.maximum(1, np.rint(bundle.units_model.predict(features))).astype(int)
    ml_pallets = np.maximum(1, np.rint(bundle.pallets_model.predict(features))).astype(int)
    ml_ldm = np.maximum(0, bundle.ldm_model.predict(features))
    ml_review = bundle.review_model.predict(features)

    comparisons = []
    pallets = []
    for pos, (_, row) in enumerate(normalized.iterrows()):
        rule = calculate_rule_result(row)
        comparisons.append({
            "Item": row["Item"],
            "Type": row["Type"],
            "Configuration": row["Configuration"],
            "Qty": int(row["Qty"]),
            "Rule units/pallet": rule["Rule units/pallet"],
            "ML units/pallet": int(ml_units[pos]),
            "Rule pallets": rule["Rule pallets"],
            "ML pallets": int(ml_pallets[pos]),
            "Rule pallet length (mm)": rule["Rule pallet length (mm)"],
            "Rule total LDM": rule["Rule total LDM"],
            "ML total LDM": round(float(ml_ldm[pos]), 3),
            "Rule manual review": rule["Rule manual review"],
            "ML manual review": str(ml_review[pos]),
            "Final result": "Manual review" if rule["Rule manual review"] == "Yes" else "Use rule calculation",
            "Reason": rule["Rule reason"],
        })
        pallets.extend(build_pallet_rows(row.to_dict(), rule))
    return pd.DataFrame(comparisons), pd.DataFrame(pallets)


def make_ml_import_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "ML Input"
    ws.append(INPUT_COLUMNS)
    ws.append(["W-01", "Window", "Openable", 1400, 1800, 12, 70, 55, 75, "Yes", "Glazed"])
    ws.append(["D-01", "Door", "Single", 1000, 2300, 6, 95, 65, 75, "Yes", "Glazed"])

    fill = PatternFill("solid", fgColor="18324A")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    validations = [
        ("B2:B1000", '"Window,Door"'),
        ("C2:C1000", '"Fixed,Openable,Mixed,Single,Double,Door + sidelight"'),
        ("J2:J1000", '"Yes,No"'),
        ("K2:K1000", '"Glazed,Unglazed,Without glass"'),
    ]
    for cells, formula in validations:
        validation = DataValidation(type="list", formula1=formula, showDropDown=False)
        ws.add_data_validation(validation)
        validation.sqref = cells

    widths = [18, 12, 20, 13, 13, 9, 18, 18, 22, 18, 16]
    for col, width in zip("ABCDEFGHIJK", widths):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def read_uploaded_workbook(file) -> Tuple[pd.DataFrame, str, List[str]]:
    sheets = pd.read_excel(file, sheet_name=None)
    warnings: List[str] = []
    if "ML Input" in sheets:
        return normalize_input(sheets["ML Input"].dropna(how="all")), "ML Input", warnings
    if "Constructions" in sheets:
        source = sheets["Constructions"].dropna(how="all").copy()
        supported = source["Type"].astype(str).map(normalize_type).isin({"Window", "Door"}) if "Type" in source else pd.Series(False, index=source.index)
        skipped = int((~supported).sum())
        if skipped:
            warnings.append(f"Skipped {skipped} unsupported row(s); the ML demo accepts only Window and Door")
        return normalize_input(source[supported]), "Legacy Constructions", warnings
    if "ML data" in sheets:
        return normalize_input(sheets["ML data"].dropna(how="all")), "Experimental ML data", warnings
    raise ValueError("Unknown workbook format. Use ML Input or the existing Constructions template.")


def make_result_workbook(input_df: pd.DataFrame, comparison_df: pd.DataFrame, pallet_df: pd.DataFrame, metrics: Dict[str, float]) -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        input_df.to_excel(writer, sheet_name="Input", index=False)
        comparison_df.to_excel(writer, sheet_name="Rule vs ML", index=False)
        pallet_df.to_excel(writer, sheet_name="Pallets", index=False)
        pd.DataFrame([{"Metric": k, "Value": v} for k, v in metrics.items()]).to_excel(writer, sheet_name="Model metrics", index=False)
    return output.getvalue()
